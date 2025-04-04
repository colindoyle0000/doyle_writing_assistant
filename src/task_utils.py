"""Utility functions common to the task modules."""

import logging
import os
import PyPDF2
from docx import Document
from langchain_chroma import Chroma
from langchain_openai import OpenAIEmbeddings

from src.utils_tokens import num_tokens
from src.utils_llm import (
    llm_call, 
    query_db, 
    filter_results_by_similarity, 
    query_db_and_filter, 
    get_documents_from_results,
    get_metadata_from_results
)
from src.utils_file import get_root_dir, extract_from_md, save_outputs
from src.source_excerpt import SourceExcerpt

# Set up logger
logger = logging.getLogger('writing_assistant')

def replace_prompt_placeholders(
    prompt,
    project,
    query="",
    manual_notes="",
):
    """Replace placeholders in the system prompt."""
    prompt = prompt.replace("{project_name}", project.project_name)
    prompt = prompt.replace("{query}", query)
    
    if len(manual_notes) == 0:
        prompt = prompt.replace("{manual_notes}", "")
    else:
        prompt = prompt.replace("{manual_notes}", f"Here are the notes you've already taken on this:\n\n{manual_notes}\n\n")
    
    if project.abstract is None:
        prompt = prompt.replace("{abstract}", "")
    else:
        abstract = f"Our abstract is\n\n{project.abstract}\n\n"
        prompt = prompt.replace("{abstract}", abstract)

    if project.structure is None:
        prompt = prompt.replace("{structure}", "")
    else:
        structure = f"Our overall article structure is:\n\n{project.structure}\n\n"
        prompt = prompt.replace("{structure}", structure)
    
    return prompt

def replace_prompt_notes_sources(
    prompt,
    notes_lst,
    sources_lst,
):
    """Replace placeholders in the system prompt."""
    if len(notes_lst) == 0:
        prompt = prompt.replace("{notes}", "")
    else:
        notes = "\n\n".join(notes_lst)
        prompt = prompt.replace("{notes}", f"Excerpts from potentially relevant notes that we've taken are:\n\n{notes}\n\n")
    
    if len(sources_lst) == 0:
        prompt = prompt.replace("{sources}", "")
    else:
        sources = "\n\n".join(sources_lst)
        prompt = prompt.replace("{sources}", f"Excerpts from potentially relevant sources are:\n\n{sources}\n\n")
    
    return prompt
    
def reduce_notes_sources_for_tokens(prompt, notes_lst, sources_lst, project):
    """Reduce notes and sources to fit within token limit."""
    # Calculate initial token count
    tokens_sum = num_tokens(prompt)
    tokens_sum += sum(num_tokens(note) for note in notes_lst)
    tokens_sum += sum(num_tokens(source) for source in sources_lst)    
    initial_notes = len(notes_lst)
    initial_sources = len(sources_lst)

    # Reduce until under token limit
    while tokens_sum > project.llm_settings.max_tokens:
        if len(notes_lst) > len(sources_lst) and notes_lst:
            note = notes_lst.pop()
            tokens_sum -= num_tokens(note)
        elif sources_lst:
            source = sources_lst.pop()
            tokens_sum -= num_tokens(source)
        else:
            break
    
    if initial_notes > len(notes_lst):
        logger.info(f"Reduced notes from {initial_notes} to {len(notes_lst)}")
    if initial_sources > len(sources_lst):
        logger.info(f"Reduced sources from {initial_sources} to {len(sources_lst)}")
    
    return notes_lst, sources_lst

def clean_up_prompt(
    prompt,
    project,
    query,
    manual_notes,
    notes_lst,
    sources_lst,
):
    """Run all prompt cleaning functions (they precede this function within this module)."""
    prompt = replace_prompt_placeholders(
        prompt=prompt,
        project=project,
        query=query,
        manual_notes=manual_notes,
    )

    notes_lst, sources_lst = reduce_notes_sources_for_tokens(
        prompt=prompt,
        notes_lst=notes_lst,
        sources_lst=sources_lst,
        project=project,
    )
    
    prompt = replace_prompt_notes_sources(
        prompt=prompt,
        notes_lst=notes_lst,
        sources_lst=sources_lst,
    )
    
    return prompt, notes_lst, sources_lst

def reduce_excerpt(
    excerpt_text: str,
    query: str,
    project,
):
    """Call on llm_cheap to copy from an excerpt only the information relevant to the query."""
    logger.info(f"Running reduce_excerpt.")
    # Build up the system prompt
    prompt_system = extract_from_md(
        os.path.join(get_root_dir(), "data", "prompts", "reduce_excerpt.md")
    )
    prompt_system += "\n\n"
    # Replace {excerpt_text} placeholder with the excerpt text
    prompt_system = prompt_system.replace("{excerpt_text}", excerpt_text)
    notes_lst = [] # Need to define these variables before calling clean_up_prompt
    sources_lst = []
    prompt_system, notes_lst, sources_lst = clean_up_prompt(
        prompt=prompt_system,
        project=project,
        query=query,
        manual_notes="",
        notes_lst=notes_lst,
        sources_lst=sources_lst,
    )
    # Create user prompt
    prompt_user = "Copy relevant information from the excerpt."
    # Set up prompts for the LLM
    prompts = [
        {"role": "system", "content": prompt_system},
        {"role": "user", "content": prompt_user},
    ]
    # Call the LLM
    response = llm_call(
        model=project.llm_settings.model_cheap,
        prompts=prompts,
        settings=project.llm_settings,
    )

    logger.info(f"Reduced excerpt from {num_tokens(excerpt_text)} to {num_tokens(response.output)} tokens.")
    return response

def reduce_excerpts_lst(
    excerpts_lst,
    query: str,
    project,
):
    """Reduce a list of excerpts to only the information relevant to the query."""
    reduced_excerpts = []
    for excerpt in excerpts_lst:
        response = reduce_excerpt(
            excerpt_text=excerpt,
            query=query,
            project=project,
        )
        reduced_excerpts.append(response.output)
    return reduced_excerpts

def retrieve_notes(
    query: str,
    project,
    k = 4,
    similarity_threshold = None,
):
    """Retrieve notes from vector_db."""
    logger.info(f"Running retrieve_notes for query: {query}")
    results_dict = query_db(
        db=project.notes.notes_db,
        query_text=query,
        num_results=k,
    )
    if similarity_threshold is not None:
        results_dict = filter_results_by_similarity(
            results=results_dict,
            similarity_threshold=similarity_threshold,
        )
    metadata_lst = get_metadata_from_results(
        results=results_dict,
    )
    
    notes_lst = get_documents_from_results(
        results=results_dict,
    )
    logger.info(f"Number of notes retrieved: {len(notes_lst)}")
    # Create excerpts list of just the text
    #logger.info(f"Reducing notes to only contain relevant information.")
    #excerpts_lst = [note for note in notes_lst]
    #notes_lst = reduce_excerpts_lst(
    #    excerpts_lst=excerpts_lst,
    #    query=query,
    #    project=project,
    #)
    return notes_lst, metadata_lst, results_dict

def retrieve_sources(
    query: str,
    project,
    k = 4,
    similarity_threshold = None,
):
    """Retrieve excerpts from vector_db."""
    logger.info(f"Running retrieve_sources for query: {query}")
    
    results_dict = query_db(
        db=project.sources.sources_db,
        query_text=query,
        num_results=k,
    )
    if similarity_threshold is not None:
        results_dict = filter_results_by_similarity(
            results=results_dict,
            similarity_threshold=similarity_threshold,
        )
    
    metadata_lst = get_metadata_from_results(
        results=results_dict,
    )
    
    sources_lst = get_documents_from_results(
        results=results_dict,
    )
    
    logger.info(f"Number of sources retrieved: {len(sources_lst)}")
    # Create excerpts list of just the text
    #logger.info(f"Reducing sources to only contain relevant information.")
    #excerpts_lst = [source for source in sources_lst]
    #sources_lst = reduce_excerpts_lst(
    #    excerpts_lst=excerpts_lst,
    #    query=query,
    #    project=project,
    #)
    # Add metadata to sources_lst
    for i, metadata_entry in enumerate(metadata_lst):
        metadata = get_excerpt_metadata(
            metadata_entry=metadata_entry,
            project=project,
        )
        sources_lst[i] = metadata + sources_lst[i]
    return sources_lst, metadata_lst, results_dict

def hypo_retrieval(
    query: str,
    project,
    num_notes: int=0,
    num_sources: int=0,
    similarity_threshold = None,
):
    """Call llm_cheap to write a hypothetical text and then retrieve similar notes and sources."""
    logger.info(f"Running hypo_retrieval for query: {query}")
    # Build up the system prompt
    prompt_system = extract_from_md(
        os.path.join(get_root_dir(), "data", "prompts", "preface.md")
    )
    prompt_system += "\n\n"
    prompt_system += extract_from_md(
        os.path.join(get_root_dir(), "data", "prompts", "hypo_retrieval.md")
    )
    notes_lst = [] # Need to define these variables before calling clean_up_prompt
    sources_lst = []
    prompt_system, notes_lst, sources_lst = clean_up_prompt(
        prompt=prompt_system,
        project=project,
        query=query,
        manual_notes="",
        notes_lst=notes_lst,
        sources_lst=sources_lst,
    )
    # Create user prompt
    prompt_user = "Write this."
    # Set up prompts for the LLM
    prompts = [
        {"role": "system", "content": prompt_system},
        {"role": "user", "content": prompt_user},
    ]
    # Call the LLM
    hypo_response = llm_call(
        model=project.llm_settings.model_cheap,
        prompts=prompts,
        settings=project.llm_settings,
    )
    # Retrieve notes
    if num_notes > 0:
        notes_lst, metadata_lst, results = retrieve_notes(
            query=hypo_response.output,
            project=project,
            k=num_notes,
            similarity_threshold=similarity_threshold,
        )
    else:
        notes_lst = []
        metadata_lst = []
        results = {}
    notes_tuple = (notes_lst, metadata_lst, results)
    # Retrieve sources
    if num_sources > 0:
        sources_lst, metadata_lst, results = retrieve_sources(
            query=hypo_response.output,
            project=project,
            k=num_sources,
            similarity_threshold=similarity_threshold,
        )
    else:
        sources_lst = []
        metadata_lst = []
        results = {}
    sources_tuple = (sources_lst, metadata_lst, results)
    save_outputs(
        directory=project.output_directory,
        filename="hypo_retrieval",
        responses=[hypo_response],
        meta_info=f"Query: {query} \n",
    )
        
    return hypo_response, notes_tuple, sources_tuple


def get_excerpt_metadata(
    metadata_entry: list,
    project,
):
    """Retrieve excerpt metadata from metadata list."""
    citekey = metadata_entry['citekey']
    excerpt_index = metadata_entry['excerpt_index']
    
    source = project.sources.sources_dict[citekey]
    excerpt = source.excerpts[excerpt_index]
    
    # If the adjusted page number is not set, set it
    if excerpt.adjusted_page_number is None:
        excerpt.get_page_number()
        excerpt.get_adjusted_page_number()
    page_number = excerpt.adjusted_page_number

    # Try to get title and author from source.fields dictionary, but default to empty string
    title = excerpt.source.fields.get('title', "")
    author = excerpt.source.fields.get('author', "")
    metadata = f"##Citekey: {citekey} \nTitle: {title} \nAuthor: {author} \nPage: {page_number} \n\n"
    return metadata

def get_full_source_from_excerpt(
    excerpt_text: str,
    project,
):
    """Retrieve source from excerpt."""
    pass

def extract_ai_notes(
    ai_notes_lst,
    project,
):
    """Extract the text of llm_responses from an ai_notes_lst comprised of a tuple of llm_responses, notes_lst, and sources_lst."""
    ai_notes_text = ""
    for ai_notes in ai_notes_lst:
        ai_notes_text += ai_notes[0].output
        ai_notes_text += "\n\n"
    return ai_notes_text

def extract_ai_notes_notes_lst(
    ai_notes_lst,
):
    """Extract the notes_lst from an ai_notes_lst comprised of a tuple of llm_responses, notes_lst, and sources_lst."""
    notes_lst = []
    for ai_notes in ai_notes_lst:
        notes_lst.extend(ai_notes[1])
    return notes_lst

def extract_ai_notes_sources_lst(
    ai_notes_lst,
):
    """Extract the sources_lst from an ai_notes_lst comprised of a tuple of llm_responses, notes_lst, and sources_lst."""
    sources_lst = []
    for ai_notes in ai_notes_lst:
        sources_lst.extend(ai_notes[2])
    return sources_lst